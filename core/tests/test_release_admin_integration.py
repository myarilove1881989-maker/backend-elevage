"""Exercise the historical admin against production, egg and birth data."""

import csv
from datetime import timedelta
from decimal import Decimal
from io import BytesIO, StringIO
from tempfile import TemporaryDirectory

from django.contrib import admin
from django.contrib.auth.models import Permission
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from core.egg_services import get_dated_egg_stock
from core.models import (
    Achat, CategorieDepense, Client, Depense, Espece, Exploitation,
    Lettrage, Lot, Mouvement, Payment, Task, User, Vente, VenteOeufs,
)


class ReleaseAdminIntegrationTests(TestCase):
    historical_models = (
        User, Exploitation, Espece, Lot, Achat, Mouvement, Vente,
        CategorieDepense, Depense, Client, Payment, Lettrage, Task,
    )

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        # Render admin pages with the real manifest storage, without writing
        # generated assets into the repository or requiring a previous build.
        assets = TemporaryDirectory(prefix="elevage-release-static-")
        cls.addClassCleanup(assets.cleanup)
        static_settings = override_settings(STATIC_ROOT=assets.name)
        static_settings.enable()
        cls.addClassCleanup(static_settings.disable)
        call_command("collectstatic", interactive=False, verbosity=0)

    @classmethod
    def setUpTestData(cls):
        cls.root = User.objects.create_superuser(username="release-root", password=None)
        cls.owner = User.objects.create_user(username="release-owner")
        cls.other = User.objects.create_user(username="release-other")
        cls.staff = User.objects.create_user(
            username="release-staff", is_staff=True, exploitation=cls.owner.exploitation,
        )
        # Ordinary Django permissions must not bypass the global-admin restriction.
        cls.staff.user_permissions.set(Permission.objects.all())
        cls.today = timezone.localdate()
        cls.species = Espece.objects.create(nom="Poulet", exploitation=cls.owner.exploitation)
        cls.parent = Lot.objects.create(
            nom="Parents release", espece=cls.species, exploitation=cls.owner.exploitation,
            date_debut=cls.today - timedelta(days=30), type_production="REPRODUCTION",
        )
        cls.customer = Client.objects.create(
            nom="Client exploitation A", exploitation=cls.owner.exploitation,
            pays="CM", ville="Yaoundé",
        )
        cls.foreign_customer = Client.objects.create(
            nom="Client exploitation B", exploitation=cls.other.exploitation,
        )
        foreign_species = Espece.objects.create(nom="Porc", exploitation=cls.other.exploitation)
        cls.foreign_lot = Lot.objects.create(
            nom="Lot exploitation B", espece=foreign_species,
            exploitation=cls.other.exploitation, date_debut=cls.today,
        )
        Achat.objects.create(
            lot=cls.parent, exploitation=cls.owner.exploitation, date=cls.today,
            quantite=100, prix_unitaire=Decimal("10"), prix_total=Decimal("1000"),
        )
        Mouvement.objects.create(
            lot=cls.parent, exploitation=cls.owner.exploitation,
            type_mouvement="ACHAT", quantite=100,
        )
        cls.animal_sale = Vente.objects.create(
            lot=cls.parent, client=cls.customer, date=cls.today,
            quantite=5, prix_unitaire=Decimal("100"),
        )
        Mouvement.objects.create(
            lot=cls.parent, exploitation=cls.owner.exploitation,
            type_mouvement="VENTE", quantite=5, client=cls.customer,
            vente=cls.animal_sale, prix_unitaire=Decimal("100"),
        )
        foreign_sale = Vente.objects.create(
            lot=cls.foreign_lot, client=cls.foreign_customer, date=cls.today,
            quantite=1, prix_unitaire=Decimal("200"),
        )
        for owner, customer, sale, amount in (
            (cls.owner, cls.customer, cls.animal_sale, 125),
            (cls.other, cls.foreign_customer, foreign_sale, 50),
        ):
            payment = Payment.objects.create(
                client=customer, exploitation=owner.exploitation,
                date=cls.today, montant=amount,
            )
            Lettrage.objects.create(vente=sale, payment=payment, montant=amount)
        category = CategorieDepense.objects.create(
            nom="Transport", exploitation=cls.owner.exploitation,
        )
        Depense.objects.create(lot=cls.parent, categorie=category, montant=Decimal("25"))
        Task.objects.create(exploitation=cls.owner.exploitation, title="Suivi release", date=cls.today)
        cls.eggs = Lot.objects.create(
            nom="Pondeuses release", espece=cls.species, exploitation=cls.owner.exploitation,
            date_debut=cls.today - timedelta(days=30), type_production="OEUFS",
        )
        api = APIClient()
        api.force_authenticate(cls.owner)
        collection = api.post("/api/oeufs/collectes/", {
            "lot": cls.eggs.pk, "nombre_collecte": 100,
            "collecte_at": (timezone.now() - timedelta(days=1)).isoformat(),
        }, format="json")
        if collection.status_code != 201:
            raise AssertionError(collection.data)
        egg_sale = api.post("/api/oeufs/ventes/", {
            "lot": cls.eggs.pk, "client": cls.customer.pk, "date": cls.today.isoformat(),
            "conditionnement": "COMPOSE", "nombre_alveoles": 1,
            "oeufs_supplementaires": 0, "prix_total": "300.00",
        }, format="json")
        if egg_sale.status_code != 201:
            raise AssertionError(egg_sale.data)
        cls.egg_sale = VenteOeufs.objects.get(pk=egg_sale.data["id"])
        birth = api.post("/api/mouvements/create/", {
            "lot": cls.parent.pk, "type_mouvement": "NAISSANCE",
            "date": cls.today.isoformat(), "total_naissances": 12, "mort_nes": 2,
            "nom_nouveau_lot": "Nouveau lot release",
        }, format="json")
        if birth.status_code != 201:
            raise AssertionError(birth.data)
        cls.child = Lot.objects.get(pk=birth.data["nouveau_lot"]["id"])

    def workbook(self):
        self.client.force_login(self.root)
        response = self.client.get(reverse("admin-export-statistics"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"],
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertRegex(response["Content-Disposition"], r'statistiques-elevage-\d{4}-\d{2}-\d{2}\.xlsx')
        return load_workbook(BytesIO(response.content))

    def test_anonymous_and_normal_users_cannot_open_global_views(self):
        urls = [reverse("superadmin-dashboard"), reverse("admin-export-statistics")]
        for user in (None, self.owner):
            self.client.logout()
            if user is not None:
                self.client.force_login(user)
            for url in urls:
                with self.subTest(user=user, url=url):
                    response = self.client.get(url)
                    self.assertEqual(response.status_code, 302)
                    self.assertIn("/admin/login/", response["Location"])

    def test_staff_with_all_model_permissions_cannot_view_or_export_global_data(self):
        self.client.force_login(self.staff)
        for name in ("superadmin-dashboard", "admin-export-statistics"):
            self.assertEqual(self.client.get(reverse(name)).status_code, 403)
        for model in self.historical_models:
            with self.subTest(model=model.__name__):
                url = reverse(f"admin:core_{model._meta.model_name}_changelist")
                self.assertEqual(self.client.get(url).status_code, 403)
                response = self.client.post(url, {
                    "action": "export_as_csv", "_selected_action": [model.objects.first().pk],
                })
                self.assertEqual(response.status_code, 403)

    def test_superadmin_dashboard_includes_egg_sales_and_payments(self):
        self.client.force_login(self.root)
        response = self.client.get(reverse("superadmin-dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, "admin/superadmin_dashboard.html")
        self.assertEqual(response.context["total_sales"], Decimal("1000"))
        self.assertEqual(response.context["total_payments"], 175)
        self.assertEqual(response.context["total_balance"], 825)
        self.assertEqual(response.context["total_expenses"], Decimal("25"))
        self.assertContains(response, reverse("admin-export-statistics"))
        self.assertEqual(len(response.context["modules"]), 12)
        for _, _, name in response.context["modules"]:
            self.assertContains(response, reverse(name))

    def test_all_historical_admins_render_and_export_populated_rows(self):
        self.client.force_login(self.root)
        for model in self.historical_models:
            with self.subTest(model=model.__name__):
                self.assertIn(model, admin.site._registry)
                url = reverse(f"admin:core_{model._meta.model_name}_changelist")
                self.assertEqual(self.client.get(url).status_code, 200)
                response = self.client.post(url, {
                    "action": "export_as_csv", "_selected_action": [model.objects.first().pk],
                })
                self.assertEqual(response.status_code, 200)
                self.assertEqual(response["Content-Type"], "text/csv; charset=utf-8")
                self.assertIn(f'{model._meta.model_name}.csv', response["Content-Disposition"])
                rows = list(csv.reader(StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
                self.assertEqual(len(rows), 2)
                self.assertEqual(len(rows[0]), len(rows[1]))
                self.assertNotIn("password", [value.lower() for value in rows[0]])

    def test_xlsx_preserves_sheets_and_financial_values_with_egg_sales(self):
        workbook = self.workbook()
        self.assertEqual(workbook.sheetnames, [
            "Exploitations", "Utilisateurs", "Espèces", "Lots", "Achats", "Mouvements",
            "Ventes", "Clients", "Paiements", "Lettrages", "Dépenses", "Catégories dépenses", "Tâches",
        ])
        sales = {row[0]: row for row in workbook["Ventes"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(sales[self.animal_sale.pk][8:12], (500, 125, 375, "PARTIEL"))
        self.assertEqual(sales[self.egg_sale.vente_id][8:12], (300, 0, 300, "IMPAYE"))
        customers = {row[0]: row for row in workbook["Clients"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(customers[self.customer.pk][6:9], (800, 125, 675))
        self.assertEqual(customers[self.foreign_customer.pk][6:9], (200, 50, 150))
        for sheet in workbook:
            self.assertGreater(sheet.max_row, 1)
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertTrue(sheet.auto_filter.ref)
        workbook.close()

    def test_csv_selection_and_calculated_customer_balance(self):
        self.client.force_login(self.root)
        response = self.client.post(reverse("admin:core_client_changelist"), {
            "action": "export_as_csv", "_selected_action": [self.customer.pk],
        })
        self.assertEqual(response.status_code, 200)
        rows = list(csv.reader(StringIO(response.content.decode("utf-8-sig")), delimiter=";"))
        self.assertEqual(len(rows), 2)
        self.assertEqual(rows[1][0], self.customer.nom)
        self.assertEqual([float(value) for value in rows[1][-3:]], [800, 125, 675])
        self.assertNotContains(response, self.foreign_customer.nom)

    def test_api_remains_scoped_even_for_staff_with_admin_permissions(self):
        api = APIClient()
        for user in (self.owner, self.staff):
            with self.subTest(user=user.username):
                api.force_authenticate(user)
                response = api.get("/api/clients/")
                self.assertEqual(response.status_code, 200)
                self.assertEqual({row["id"] for row in response.data}, {self.customer.pk})
                self.assertIn(api.get(f"/api/lots/{self.foreign_lot.pk}/").status_code, (403, 404))

    def test_exports_preserve_birth_stock_and_do_not_change_fifo_allocations(self):
        before = get_dated_egg_stock(self.owner.exploitation, self.eggs)
        self.assertEqual(before["stock_global"], 70)
        self.assertTrue(before["origines_completes"])
        self.assertEqual(self.child.mouvements.get().lot_origine_id, self.parent.pk)
        self.assertEqual((self.parent.stock, self.child.stock), (95, 10))
        workbook = self.workbook()
        stocks = {row[0]: row[4] for row in workbook["Lots"].iter_rows(min_row=2, values_only=True)}
        self.assertEqual((stocks[self.parent.pk], stocks[self.child.pk], stocks[self.eggs.pk]), (95, 10, 0))
        self.assertEqual(get_dated_egg_stock(self.owner.exploitation, self.eggs), before)
        workbook.close()
