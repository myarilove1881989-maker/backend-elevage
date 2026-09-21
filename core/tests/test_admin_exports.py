from io import BytesIO

from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from core.models import User


class AdminExportTests(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser(
            username="export-admin",
            email="admin@example.com",
            password="TestPass123!",
        )
        self.client.force_login(self.user)

    def test_global_excel_export_contains_expected_sheets(self):
        response = self.client.get(reverse("admin-export-statistics"))
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        self.assertIn("Clients", workbook.sheetnames)
        self.assertIn("Paiements", workbook.sheetnames)
        self.assertIn("Lettrages", workbook.sheetnames)
        self.assertIn("Ventes", workbook.sheetnames)

    def test_csv_export_action_is_available(self):
        model_admin = __import__("core.admin", fromlist=["ClientAdmin"]).ClientAdmin
        self.assertIn("export_as_csv", model_admin.actions)
