from io import BytesIO

from django.contrib.admin.views.decorators import staff_member_required
from django.core.exceptions import PermissionDenied
from django.db.models import Sum
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import (
    Achat, CategorieDepense, Client, Depense, Espece, Exploitation,
    Lettrage, Lot, Mouvement, Payment, Task, User, Vente,
)


def text(value):
    return "" if value is None else str(value)


def date_value(value):
    if value is None:
        return ""
    if hasattr(value, "tzinfo") and value.tzinfo is not None:
        value = timezone.localtime(value)
    return value.replace(tzinfo=None) if hasattr(value, "replace") and hasattr(value, "tzinfo") else value


def add_sheet(workbook, title, headers, rows):
    sheet = workbook.create_sheet(title=title)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="073B5C")
    sheet.freeze_panes = "A2"
    for row in rows:
        sheet.append(list(row))

    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"

    for index, column in enumerate(sheet.columns, 1):
        longest = max((len(text(cell.value)) for cell in column), default=8)
        sheet.column_dimensions[get_column_letter(index)].width = min(max(longest + 2, 12), 42)


@staff_member_required
def export_all_statistics(request):
    if not request.user.is_superuser:
        raise PermissionDenied

    workbook = Workbook()
    workbook.remove(workbook.active)

    add_sheet(workbook, "Exploitations",
        ["ID", "Nom", "Propriétaire", "Utilisateurs", "Lots", "Clients", "CA FCFA", "Paiements FCFA", "Solde clients FCFA", "Dépenses FCFA", "Création"],
        ((farm.id, farm.nom, farm.proprietaire.username, farm.users.count(), farm.lots.count(), farm.clients.count(),
          float(Vente.objects.filter(lot__exploitation=farm).aggregate(v=Sum("montant_total"))["v"] or 0),
          float(Payment.objects.filter(exploitation=farm).aggregate(v=Sum("montant"))["v"] or 0),
          float(Vente.objects.filter(lot__exploitation=farm).aggregate(v=Sum("montant_total"))["v"] or 0) - float(Payment.objects.filter(exploitation=farm).aggregate(v=Sum("montant"))["v"] or 0),
          float(Depense.objects.filter(lot__exploitation=farm).aggregate(v=Sum("montant"))["v"] or 0), date_value(farm.date_creation))
         for farm in Exploitation.objects.select_related("proprietaire").all()))

    add_sheet(workbook, "Utilisateurs",
        ["ID", "Utilisateur", "E-mail", "Prénom", "Nom", "Exploitation", "Actif", "Superutilisateur", "Inscription", "Dernière connexion"],
        ((u.id, u.username, u.email, u.first_name, u.last_name, text(u.exploitation), u.is_active, u.is_superuser, date_value(u.date_joined), date_value(u.last_login))
         for u in User.objects.select_related("exploitation").all()))

    add_sheet(workbook, "Espèces", ["ID", "Nom", "Exploitation"],
        ((x.id, x.nom, text(x.exploitation)) for x in Espece.objects.select_related("exploitation").all()))

    add_sheet(workbook, "Lots", ["ID", "Nom", "Espèce", "Exploitation", "Stock", "Date début", "Date fin", "Prix vente prévu"],
        ((x.id, x.nom, x.espece.nom, x.exploitation.nom, x.stock, x.date_debut, x.date_fin, x.prix_vente_prevu)
         for x in Lot.objects.select_related("espece", "exploitation").all()))

    add_sheet(workbook, "Achats", ["ID", "Date", "Exploitation", "Lot", "Quantité", "Prix unitaire", "Prix total", "Fournisseur", "Note", "Créé par"],
        ((x.id, x.date, x.exploitation.nom, x.lot.nom, x.quantite, x.prix_unitaire, x.prix_total, text(x.fournisseur), text(x.note), text(x.created_by))
         for x in Achat.objects.select_related("exploitation", "lot", "created_by").all()))

    add_sheet(workbook, "Mouvements", ["ID", "Date", "Exploitation", "Lot", "Espèce", "Type", "Client", "Quantité", "Prix unitaire", "Montant total", "Créé par"],
        ((x.id, x.date, text(x.exploitation), x.lot.nom, x.lot.espece.nom, x.type_mouvement, text(x.client), x.quantite, x.prix_unitaire, x.montant_total, text(x.created_by))
         for x in Mouvement.objects.select_related("exploitation", "lot__espece", "client", "created_by").all()))

    add_sheet(workbook, "Ventes", ["ID", "Date", "Exploitation", "Client", "Lot", "Espèce", "Quantité", "Prix unitaire", "Facturé", "Payé", "Reste", "Statut"],
        ((x.id, x.date, x.lot.exploitation.nom, text(x.client), x.lot.nom, x.lot.espece.nom, x.quantite, x.prix_unitaire, x.montant_total, x.montant_paye, x.reste_a_payer, x.statut)
         for x in Vente.objects.select_related("lot__exploitation", "lot__espece", "client").all()))

    add_sheet(workbook, "Clients", ["ID", "Nom", "Téléphone", "Pays", "Ville", "Exploitation", "Facturé", "Payé", "Solde"],
        ((x.id, x.nom, x.telephone, x.pays, x.ville, text(x.exploitation),
          float(x.ventes.aggregate(v=Sum("montant_total"))["v"] or 0), float(x.payments.aggregate(v=Sum("montant"))["v"] or 0),
          float(x.ventes.aggregate(v=Sum("montant_total"))["v"] or 0) - float(x.payments.aggregate(v=Sum("montant"))["v"] or 0))
         for x in Client.objects.select_related("exploitation").all()))

    add_sheet(workbook, "Paiements", ["ID", "Date", "Exploitation", "Client", "Montant", "Montant lettré", "Non affecté", "Note", "Création"],
        ((x.id, x.date, x.exploitation.nom, x.client.nom, x.montant,
          float(x.lettrages.aggregate(v=Sum("montant"))["v"] or 0), x.montant - float(x.lettrages.aggregate(v=Sum("montant"))["v"] or 0), text(x.note), date_value(x.created_at))
         for x in Payment.objects.select_related("exploitation", "client").all()))

    add_sheet(workbook, "Lettrages", ["ID", "Date", "Exploitation", "Client", "Vente ID", "Paiement ID", "Montant affecté"],
        ((x.id, date_value(x.created_at), x.vente.lot.exploitation.nom, text(x.vente.client or x.payment.client), x.vente_id, x.payment_id, x.montant)
         for x in Lettrage.objects.select_related("vente__lot__exploitation", "vente__client", "payment__client").all()))

    add_sheet(workbook, "Dépenses", ["ID", "Date", "Exploitation", "Lot", "Espèce", "Catégorie", "Montant", "Note"],
        ((x.id, x.date, x.lot.exploitation.nom, x.lot.nom, x.lot.espece.nom, text(x.categorie), x.montant, text(x.note))
         for x in Depense.objects.select_related("lot__exploitation", "lot__espece", "categorie").all()))

    add_sheet(workbook, "Catégories dépenses", ["ID", "Nom", "Exploitation"],
        ((x.id, x.nom, x.exploitation.nom) for x in CategorieDepense.objects.select_related("exploitation").all()))

    add_sheet(workbook, "Tâches", ["ID", "Date", "Exploitation", "Titre", "Création"],
        ((x.id, x.date, x.exploitation.nom, x.title, date_value(x.created_at)) for x in Task.objects.select_related("exploitation").all()))

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    filename = f"statistiques-elevage-{timezone.localdate().isoformat()}.xlsx"
    response = HttpResponse(output.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response
